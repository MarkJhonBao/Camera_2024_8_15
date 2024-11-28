import torch, math, time
import torch.nn as nn
from tqdm import *
import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference


def GetLOSS(LOSS_List, LossWeight):
    LOSS = LossWeight[0] * LOSS_List[0]
    LOSS_value = LOSS_List[0].item()
    for ii in range(len(LOSS_List) - 1):
        LOSS = LOSS + LossWeight[ii + 1] * LOSS_List[ii + 1]
        LOSS_value = LOSS_value + LOSS_List[ii + 1].item()
    return LOSS, LOSS_value


def UpdateLossWeight(model, LOSS_List, LossWeight, beta=0.9):
    Last_LossWeight = torch.tensor(LossWeight).clone().detach()
    L1 = LOSS_List[0]
    model.zero_grad()
    tgrad = ObtainGrad(model, L1)
    max_value = torch.max(torch.abs(tgrad))
    mean_value = []
    for ii in range(len(LossWeight) - 1):
        model.zero_grad()
        tL = LOSS_List[ii + 1]
        try:
            tgrad = ObtainGrad(model, tL)
            tmean_value = torch.mean(torch.abs(tgrad))
        except:
            tmean_value = 0
        mean_value.append(tmean_value)
        if tmean_value == 0:
            tLossWeight = 0
        else:
            tLossWeight = 1 / LossWeight[ii + 1] * max_value / tmean_value
        LossWeight[ii + 1] = beta * LossWeight[ii + 1] + (1 - beta) * tLossWeight
    if torch.isinf(torch.tensor(LossWeight)).any() or torch.isnan(torch.tensor(LossWeight)).any():
        return Last_LossWeight
    return torch.tensor(LossWeight)


def ObtainGrad(model, L, retain_graph=True):
    L.backward(retain_graph=retain_graph)
    ii = 0
    for name, parms in model.named_parameters():
        tgrad = parms.grad
        if len(tgrad.size()) == 1:
            size = tgrad.size()[0]
        else:
            size = tgrad.size()[0] * tgrad.size()[1]
        if ii == 0:
            grad_value = tgrad.view(size, 1)
            ii += 1
            continue
        tgrad = tgrad.view(size, 1)
        grad_value = torch.cat((grad_value, tgrad))
    return grad_value


def Integration1D(U, X, LB, UB):
    TempX = X.squeeze()
    TempU = U.squeeze()
    TempX, Index = torch.sort(TempX)
    TempU = TempU[Index]
    Index1 = TempX >= LB
    Index2 = TempX <= UB
    Index = Index1 * Index2
    TempX = TempX[Index]
    TempU = TempU[Index]
    Integral = torch.trapz(TempU, TempX)
    return Integral


class BaseNetwork(nn.Module):
    def __init__(self, act_fn, input_size=3, output_size=2, hidden_sizes=[32, 32, 32], anainfo=[]):
        super().__init__()
        layers = []
        layer_sizes = [input_size] + hidden_sizes
        if type(act_fn) != type([]):
            tact_fn = act_fn
            act_fn = []
            for ii in range(len(layer_sizes) - 1):
                act_fn.append(tact_fn)
        for layer_index in range(1, len(layer_sizes)):
            layers += [nn.Linear(layer_sizes[layer_index - 1], layer_sizes[layer_index]),
                       act_fn[layer_index - 1]]
        layers += [nn.Linear(layer_sizes[-1], output_size)]
        self.layers = nn.Sequential(*layers)
        act_fn_name = []
        for ii in act_fn:
            act_fn_name.append(ii._get_name())
        self.config = {"act_fn": act_fn_name, "input_size": input_size, "output_size": output_size,
                       "hidden_sizes": hidden_sizes}

    def forward(self, x):
        out = self.layers(x)
        return out


class LossValue:
    Loss_ubar = []
    Loss_theta = []
    Loss_BC1 = []
    Loss_BC2 = []
    #Loss_BC3 = []
    # Loss_BC4 = []
    #Loss_BC5 = []


def gradients(U, X, order=1):
    if order == 1:
        return torch.autograd.grad(U, X, grad_outputs=torch.ones_like(U),
                                   create_graph=True,
                                   retain_graph=True,
                                   only_inputs=True, )[0]
    else:
        return gradients(gradients(U, X), X, order=order - 1)


def get_global_dis(S, U0, Theta, L):
    U = torch.zeros_like(S)
    V = torch.zeros_like(S)
    dU = -(1 - torch.cos(Theta)) + torch.cos(Theta) * gradients(U0, S, 1) / L
    dV = torch.sin(Theta) * (1 + gradients(U0, S, 1) / L)
    for ii in range(len(S)):
        U[ii] = Integration1D(dU, S, 0, S[ii]) * L
        V[ii] = Integration1D(dV, S, 0, S[ii]) * L
    return U, V


def get_soilpressure(S, V, k1, k2, L):
    p = -(k1 + (k2 - k1) * S) * V
    return p


def get_axialforce(S, Theta, Fx1, Fy1, p, L):
    P = Fx1 * torch.cos(Theta) + Fy1 * torch.sin(Theta)
    for ii in range(len(P)):
        P[ii] = P[ii] + Integration1D(p, S, S[ii], 1) * L * torch.sin(Theta[ii])
    return P


def get_bendingmoment(S, U, V, Fx1, Fy1, M1, p, L):
    U_S1 = U[-1]
    V_S1 = V[-1]
    M = M1 + Fy1 * (L + U_S1 - S * L - U) - Fx1 * (V_S1 - V)
    for ii in range(len(M)):
        tU = U[ii]
        tS = S[ii]
        dM = p * (S * L + U - tS * L - tU)
        M[ii] = M[ii] + Integration1D(dM, S, S[ii], 1) * L
    return M


def train_model(model, anainfo, model_path, model_name, num_sample=50, num_epochs=5000, TOL=1E-7, lr=0.001, alw=True):
    model.train()
    opt = torch.optim.Adam(params=model.parameters(), lr=lr)
    loss = torch.nn.MSELoss(reduction="sum")

    E, I, A, L = anainfo["E"], anainfo["I"], anainfo["A"], anainfo["L"]
    k1, k2 = anainfo["k1"], anainfo["k2"]
    Fx1, Fy1, M1 = anainfo["Fx1"], anainfo["Fy1"], anainfo["M1"]

    LossWeight = [1, 1, 1, 1]
    desc = "start training..."
    pbar = tqdm(range(num_epochs), desc=desc)
    MinLOSS = 100

    S = torch.rand(num_sample - 2, 1)
    S, S_indice = torch.sort(S, dim=0)
    S = torch.cat((S, torch.tensor([[1]])))
    S = torch.cat((torch.tensor([[0]]), S))
    S.requires_grad = True

    for epoch in pbar:
        Results = model(S)
        U0, Theta = Results[:, 0], Results[:, 1]
        U0 = U0.view(len(U0), 1)
        Theta = Theta.view(len(Theta), 1)

        U, V = get_global_dis(S, U0, Theta, L)
        p = get_soilpressure(S, V, k1, k2, L)
        P = get_axialforce(S, Theta, Fx1, Fy1, p, L)
        M = get_bendingmoment(S, U, V, Fx1, Fy1, M1, p, L)

        ###归一化系数
        n = num_sample
        N_theta = (1 / n) * torch.sum((M / (E * I)) ** 2)
        N_Delta = (1 / n) * torch.sum((P / (E * I)) ** 2)
        N_B_theta = (1 / n) * torch.sum(Theta ** 2)
        N_B_Delta = (1 / n) * torch.sum(U0 ** 2)
        ###桩顶无竖向约束
        # N_Tu = Fx1 ** 2+Fy1** 2+M1** 2
        ###桩顶有竖向约束
        #N_Tu = (1 / n) * torch.sum(U ** 2)
        ###桩顶无水平约束
        # N_Ty = Fx1 ** 2+Fy1** 2+M1** 2
        ###桩顶有水平约束
        # N_Ty = (1 / n) * torch.sum(V**2)
        ###桩顶无转角约束
        # N_Ttheta = Fx1 ** 2+Fy1** 2+M1** 2
        ###桩顶有转角约束
        #N_Ttheta = (1 / n) * torch.sum(Theta ** 2)

        ### 控制方程
        Cond1 = P / (E * A)
        L1 = N_Delta * loss(gradients(U0, S, 1) / L, Cond1) / torch.norm(Cond1, p=2)
        Cond2 = M / (E * I)
        L2 = N_theta * loss(gradients(Theta, S, 1) / L, Cond2) / torch.norm(Cond2, p=2)
        ### 桩底边界条件
        U0_S0, Theta_S0 = U0[0].view(1, 1), Theta[0].view(1, 1)
        Cond3 = torch.zeros(1, 1)
        L3 = N_B_Delta * loss(U0_S0, Cond3) / (torch.norm(U0, p=2) / num_sample)
        Cond4 = torch.zeros(1, 1)
        L4 = N_B_theta * loss(Theta_S0, Cond4) / (torch.norm(Theta, p=2) / num_sample)
        ### 桩顶边界条件
        U0_S1, Theta_S1, U_S1, V_S1 = U0[-1].view(1, 1), Theta[-1].view(1, 1), U[-1].view(1, 1), V[-1].view(1, 1)
        # Cond5 = torch.tensor(Fx1,dtype=torch.float32)
        # Cond6 = torch.tensor(Fy1,dtype=torch.float32)
        # Cond7 = torch.tensor(M1,dtype=torch.float32)
        #Cond5 = torch.zeros(1, 1)
        # Cond6 = torch.zeros(1,1)
        #Cond7 = torch.zeros(1, 1)

        # 桩顶无竖向约束
        # L5 = N_Tu*loss(E*A * torch.cos(U0)*gradients(U0_S1, S, 1) / L- E*I / L**2 * torch.sin(Theta) * gradients(Theta_S1, S, 2), Cond5) / torch.norm(Cond5, p=2)
        # 桩顶有竖向约束
        #L5 = N_Tu * loss(U_S1, Cond5) / (torch.norm(U_S1, p=2) / num_sample)

        # 桩顶无水平约束
        # L6 = N_Ty*loss(E*A * torch.sin(U0)*gradients(U0_S1, S, 1) / L+ E*I / L**2 * torch.cos(Theta) * gradients(Theta_S1, S, 2), Cond6) / torch.norm(Cond6, p=2)

        # 桩顶有水平约束
        # L6 = N_Ty*loss(V_S1, Cond6) / (torch.norm(V_S1, p=2) / num_sample)

        # 桩顶无转角约束
        # L7 = N_Ttheta*loss(E*I*gradients(Theta_S1, S, 1)/ L , Cond7) / torch.norm(Cond7, p=2)

        # 桩顶有转角约束
        #L7 = N_Ttheta * loss(Theta_S1, Cond7) / (torch.norm(Theta_S1, p=2) / num_sample)

        opt.zero_grad()
        LossFuncSum = [L1, L2, L3, L4]
        if epoch % 5 == 0 and epoch > 10:
            if alw:
                LossWeight = UpdateLossWeight(model, LossFuncSum, LossWeight)
            LOSS, LOSS_value = GetLOSS(LossFuncSum, LossWeight)
            opt.zero_grad()
            LOSS.backward()
            opt.step()
        else:
            LOSS, LOSS_value = GetLOSS(LossFuncSum, LossWeight)
            LOSS.backward()
            opt.step()
        if LOSS_value < TOL:
            torch.save(model, model_path + r"\\" + model_name + ".pth")
            break
        if epoch % 10 == 0 and epoch != 0 and (LOSS_value < MinLOSS):
            MinLOSS = LOSS_value
            try:
                torch.save(model, model_path + r"\\" + model_name + ".pth")
            except:
                pass

        LossValue.Loss_ubar.append(L1.item())
        LossValue.Loss_theta.append(L2.item())
        LossValue.Loss_BC1.append(L3.item())
        LossValue.Loss_BC2.append(L4.item())
        #LossValue.Loss_BC3.append(L5.item())
        # LossValue.Loss_BC4.append(L6.item())
        #LossValue.Loss_BC5.append(L7.item())

        LossResults = ""
        for ii in range(len(LossFuncSum)):
            LossResults = LossResults + '%.2e' % LossFuncSum[ii].item() + ";"
        pbar.set_description(
            "Current Loss %.2e" % LOSS_value + " Min Loss %.2e" % MinLOSS + " Loss terms:" + LossResults)
    model.eval()
    return model, epoch


def get_distop(model, num_sample):
    L = model.config["anainfo"]["L"]
    x = torch.linspace(0, 1, num_sample).view(num_sample, 1)
    x.requires_grad = True
    results = model(x)
    U0, Theta = results[:, 0].view(num_sample, 1), results[:, 1].view(num_sample, 1)
    U, V = get_global_dis(x, U0, Theta, L)
    return U[-1].item(), V[-1].item()


if __name__ == '__main__':
    Disy, Disx = {}, {}
    Epoch = {}
    Time = {}
    k = 200
    anainfo = {"E": 2e8, "I": 2.8179e-4, "A": 2.275e-2, "L": 25, "k1": k, "k2": k, "Fx1": -800, "Fy1": 150, "M1": 0}
    numLayer = 3
    numNeuron = 50
    hidden_sizes = (torch.zeros(numLayer) + int(numNeuron)).int().tolist()
    model = BaseNetwork(act_fn=[nn.Tanhshrink(), nn.Tanh(), nn.Tanhshrink()], input_size=1, output_size=2,
                        hidden_sizes=hidden_sizes)
    model.config["anainfo"] = anainfo
    model_path = os.getcwd() + r"\\model"
    model_name = "LargeDeflection"
    staT = time.time()
    model, tepoch = train_model(model=model, anainfo=anainfo, model_path=model_path, model_name=model_name, alw=True,
                                lr=0.001)

    num_sample = 50
    S = torch.linspace(0, 1, num_sample).view(num_sample, 1)
    S.requires_grad = True

    Results = model(S)
    U0, Theta = Results[:, 0].view(num_sample, 1), Results[:, 1].view(num_sample, 1)
    U, V = get_global_dis(S, U0, Theta, anainfo["L"])
    P = get_axialforce(S, Theta, anainfo["Fx1"], anainfo["Fy1"],
                       get_soilpressure(S, V, anainfo["k1"], anainfo["k2"], anainfo["L"]), anainfo["L"])
    M = get_bendingmoment(S, U, V, anainfo["Fx1"], anainfo["Fy1"], anainfo["M1"],
                          get_soilpressure(S, V, anainfo["k1"], anainfo["k2"], anainfo["L"]), anainfo["L"])

    data = {
        "采样点": S.detach().numpy().flatten() * 25,
        "位移 U": U.detach().numpy().flatten(),
        "位移 V": V.detach().numpy().flatten(),
        "轴力 P": P.detach().numpy().flatten(),
        "弯矩 M": M.detach().numpy().flatten()
    }
    df = pd.DataFrame(data)
    excel_path = model_path + "\\归一化data.e-7.xlsx"
    df.to_excel(excel_path, index=False)

    print("数据已保存至: " + excel_path)

    wb = load_workbook(excel_path)
    ws = wb.active

    chart_displacement = LineChart()
    chart_displacement.title = "位移图"
    chart_displacement.x_axis.title = "采样点"
    chart_displacement.y_axis.title = "位移"

    data_ref = Reference(ws, min_col=2, min_row=1, max_col=3, max_row=num_sample + 1)
    chart_displacement.add_data(data_ref, titles_from_data=True)
    ws.add_chart(chart_displacement, "G2")
    chart_axial_force = LineChart()
    chart_axial_force.title = "轴力图"
    chart_axial_force.x_axis.title = "采样点"
    chart_axial_force.y_axis.title = "轴力"

    data_ref = Reference(ws, min_col=4, min_row=1, max_col=4, max_row=num_sample + 1)
    chart_axial_force.add_data(data_ref, titles_from_data=True)
    ws.add_chart(chart_axial_force, "G20")

    chart_bending_moment = LineChart()
    chart_bending_moment.title = "弯矩图"
    chart_bending_moment.x_axis.title = "采样点"
    chart_bending_moment.y_axis.title = "弯矩"

    data_ref = Reference(ws, min_col=5, min_row=1, max_col=5, max_row=num_sample + 1)
    chart_bending_moment.add_data(data_ref, titles_from_data=True)
    ws.add_chart(chart_bending_moment, "G38")

    wb.save(excel_path)
    print("已在 Excel 文件中生成图表。")